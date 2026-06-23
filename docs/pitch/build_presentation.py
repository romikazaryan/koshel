#!/usr/bin/env python3
"""Generate koshel pitch deck (PPTX + PDF) — visual-first layout."""

from __future__ import annotations

import shutil
from pathlib import Path

from fpdf import FPDF
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
ASSETS = Path("/Users/romik/.cursor/projects/Users-romik-Desktop-koshel/assets")
OUT = ROOT
SCREENSHOTS = OUT / "screenshots"
PROCESSED = OUT / "processed"

SCREEN_MAP = {
    "dashboard.png": "photo_2026-06-16_12.24.26-62149caf-4aa6-4d47-9cc4-692bd8687f3e.png",
    "expenses.png": "photo_2026-06-16_12.24.49-e527da13-9a2f-44fe-a1ad-242b223f5ab5.png",
    "income.png": "photo_2026-06-16_12.24.52-084b5da1-12a7-4da2-941e-d0968002d873.png",
    "capital_overview.png": "photo_2026-06-16_12.24.18-b841db24-73f6-4fd4-8804-bda5aabbd9b3.png",
    "capital_stocks.png": "photo_2026-06-16_12.24.28-8b2e2611-8790-471b-ba7c-e734e819f8e9.png",
    "capital_crypto.png": "photo_2026-06-16_12.24.31-7d413db1-8870-41a1-817f-e438dc7c4805.png",
    "add_asset.png": "photo_2026-06-16_12.24.54-e7502faf-d641-4004-b279-26f135ecc0c1.png",
    "connections_tinvest.png": "photo_2026-06-16_12.24.40-aa293706-fbdc-454a-b9fc-05ad44f6ddd3.png",
    "connections_tbank.png": "photo_2026-06-16_12.24.36-1f004fbb-b878-4d3a-ba78-14cd1c555a52.png",
    "profile.png": "photo_2026-06-16_12.24.33-3c7b5387-b8e5-4829-be8d-f370a0ea899f.png",
}

# Colors
ACCENT = RGBColor(0x4F, 0x46, 0xE5)
DARK = RGBColor(0x0F, 0x17, 0x2A)
DARK_SOFT = RGBColor(0x1E, 0x29, 0x3B)
TEXT = RGBColor(0xF8, 0xFA, 0xFC)
TEXT_DARK = RGBColor(0x1E, 0x29, 0x3B)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
MUTED_DARK = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BG_LIGHT = RGBColor(0xF1, 0xF5, 0xF9)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Every content slide has at least one screenshot.
SLIDES: list[dict] = [
    {
        "layout": "cover",
        "title": "koshel",
        "subtitle": "Личные финансы в одном приложении",
        "tagline": "Запрос на пилот Open Finance · T‑Банк",
        "images": ["dashboard.png", "capital_overview.png", "connections_tbank.png"],
    },
    {
        "layout": "split",
        "theme": "dark",
        "title": "Проблема",
        "subtitle": "Деньги и траты разбросаны по разным приложениям",
        "bullets": [
            "Карты, вклады, брокер, крипта — нет единой картины",
            "Ручной учёт утомляет — траты не фиксируют",
            "Пользователь не видит: «сколько у меня всего»",
        ],
        "image": "capital_crypto.png",
    },
    {
        "layout": "split",
        "theme": "light",
        "title": "Решение koshel",
        "subtitle": "Расходы + доходы + весь капитал",
        "bullets": [
            "Минимум действий: голос, чек, автоимпорт из банка",
            "Живые курсы: MOEX, CoinGecko, ЦБ РФ",
            "AI-анализ и индекс финансового здоровья",
        ],
        "image": "profile.png",
    },
    {
        "layout": "hero",
        "title": "Главная",
        "subtitle": "Индекс здоровья · категории трат · AI-рекомендации",
        "image": "dashboard.png",
    },
    {
        "layout": "hero",
        "title": "Расходы",
        "subtitle": "Голос · скан чека · категории · долги и подписки",
        "image": "expenses.png",
    },
    {
        "layout": "hero",
        "title": "Доходы",
        "subtitle": "Зарплата, подработка и другие поступления",
        "image": "income.png",
    },
    {
        "layout": "hero",
        "title": "Капитал",
        "subtitle": "353 795 руб. — все активы, динамика за период",
        "image": "capital_overview.png",
    },
    {
        "layout": "hero",
        "title": "Акции · T‑Invest",
        "subtitle": "Портфель MOEX, P&L, синхронизация брокера",
        "image": "capital_stocks.png",
    },
    {
        "layout": "hero",
        "title": "Крипта",
        "subtitle": "CoinGecko, единый рублёвый итог, динамика",
        "image": "capital_crypto.png",
    },
    {
        "layout": "hero",
        "title": "Все классы активов",
        "subtitle": "Вклады · акции · облигации · крипта · недвижимость",
        "image": "add_asset.png",
    },
    {
        "layout": "dual",
        "title": "Подключения",
        "subtitle": "T‑Invest уже работает · T‑Банк — следующий шаг",
        "images": ["connections_tinvest.png", "connections_tbank.png"],
        "labels": ["T‑Invest · активно", "T‑Банк · Open Finance"],
    },
    {
        "layout": "split",
        "theme": "light",
        "title": "Ценность для T‑Банка",
        "subtitle": "Win-win, не односторонний доступ к данным",
        "bullets": [
            "Рекомендации вкладов при «свободных деньгах»",
            "Cross-sell: карты и кэшбэк под реальные категории трат",
            "Удержание: T‑Invest + счета T‑Банка в одном PFM",
            "Пилот Open Finance ЦБ с готовым продуктом",
            "Только OAuth, без паролей и парсинга",
        ],
        "image": "dashboard.png",
    },
    {
        "layout": "split",
        "theme": "dark",
        "title": "Техническая готовность",
        "subtitle": "Инфраструктура Open Finance уже в продакшене",
        "bullets": [
            "financial_connections + secrets (server-only)",
            "Edge Functions: bank-sync, tinvest-sync",
            "Дедупликация транзакций, source=bank с сервера",
            "Согласие пользователя и отзыв доступа",
        ],
        "image": "connections_tinvest.png",
    },
    {
        "layout": "closing",
        "title": "Запрос к T‑Банку",
        "subtitle": "Пилот read-only доступа к операциям физлиц",
        "bullets": [
            "OAuth-согласие на чтение транзакций",
            "Импорт 90 дней + инкрементальный синк",
            "Демо на iPhone · техдокументация · ограниченный пилот",
        ],
        "contact": "openapi@tbank.ru",
        "image": "connections_tbank.png",
    },
]


def copy_screenshots() -> None:
    SCREENSHOTS.mkdir(parents=True, exist_ok=True)
    for dest, src_name in SCREEN_MAP.items():
        src = ASSETS / src_name
        if not src.exists():
            raise FileNotFoundError(src)
        shutil.copy2(src, SCREENSHOTS / dest)


def _rounded_mask(size: tuple[int, int], radius: int) -> Image.Image:
    mask = Image.new("L", size, 0)
    draw = ImageDraw.Draw(mask)
    draw.rounded_rectangle((0, 0, size[0] - 1, size[1] - 1), radius=radius, fill=255)
    return mask


def prepare_phone_image(path: Path, max_h: int = 1200, radius: int = 48) -> Path:
    """Crop status bar, resize, round corners — for crisp slides."""
    out = PROCESSED / f"framed_{path.stem}_{max_h}.png"
    out.parent.mkdir(parents=True, exist_ok=True)
    if out.exists() and out.stat().st_mtime > path.stat().st_mtime:
        return out

    with Image.open(path) as img:
        img = img.convert("RGBA")
        w, h = img.size
        crop_top = int(h * 0.045)
        img = img.crop((0, crop_top, w, h))
        ratio = max_h / img.height
        new_w = int(img.width * ratio)
        img = img.resize((new_w, max_h), Image.Resampling.LANCZOS)
        mask = _rounded_mask(img.size, radius)
        bg = Image.new("RGBA", img.size, (0, 0, 0, 0))
        bg.paste(img, (0, 0), mask)
        bg.save(out, optimize=True)
    return out


def set_slide_bg(slide, rgb: RGBColor) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text: str,
    *,
    size: int = 24,
    bold: bool = False,
    color: RGBColor = TEXT_DARK,
    align=PP_ALIGN.LEFT,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align


def add_bullets(
    slide,
    left,
    top,
    width,
    height,
    bullets: list[str],
    *,
    size: int = 17,
    color: RGBColor = TEXT_DARK,
    spacing: int = 10,
) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(spacing)
        p.level = 0
        p.bullet = True


def place_phone(slide, img_name: str, left, top, height) -> None:
    """Phone with soft shadow."""
    img_path = prepare_phone_image(SCREENSHOTS / img_name, max_h=int(height.inches * 220))
    shadow_w = Inches(height.inches * 0.48)
    shadow = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, left + Inches(0.06), top + Inches(0.08), shadow_w, height
    )
    shadow.fill.solid()
    shadow.fill.fore_color.rgb = RGBColor(0x00, 0x00, 0x00)
    shadow.fill.transparency = 0.82
    shadow.line.fill.background()
    shadow.adjustments[0] = 0.12
    slide.shapes.add_picture(str(img_path), left, top, height=height)


def add_cover_slide(prs: Presentation, data: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK)

    # Accent glow
    glow = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.5), Inches(-1.5), Inches(6), Inches(6))
    glow.fill.solid()
    glow.fill.fore_color.rgb = ACCENT
    glow.fill.transparency = 0.88
    glow.line.fill.background()

    add_text(slide, Inches(0.9), Inches(1.0), Inches(6.5), Inches(1.2), data["title"], size=56, bold=True, color=WHITE)
    add_text(slide, Inches(0.9), Inches(2.1), Inches(6.2), Inches(0.8), data["subtitle"], size=24, color=MUTED)
    add_text(slide, Inches(0.9), Inches(2.9), Inches(6.5), Inches(0.6), data["tagline"], size=18, color=RGBColor(0xA5, 0xB4, 0xFC))

    imgs = data["images"]
    positions = [Inches(6.8), Inches(8.55), Inches(10.3)]
    heights = [Inches(5.2), Inches(4.6), Inches(4.0)]
    tops = [Inches(1.0), Inches(1.35), Inches(1.7)]
    for img, left, top, h in zip(imgs, positions, tops, heights):
        place_phone(slide, img, left, top, h)


def add_hero_slide(prs: Presentation, data: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, BG_LIGHT)

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(1.15))
    bar.fill.solid()
    bar.fill.fore_color.rgb = DARK
    bar.line.fill.background()

    add_text(slide, Inches(0.75), Inches(0.22), Inches(8), Inches(0.55), data["title"], size=30, bold=True, color=WHITE)
    add_text(slide, Inches(0.75), Inches(0.72), Inches(10), Inches(0.4), data["subtitle"], size=15, color=MUTED)

    phone_h = Inches(6.15)
    phone_w = Inches(6.15 * 0.46)
    left = (SLIDE_W - phone_w) / 2
    place_phone(slide, data["image"], left, Inches(1.22), phone_h)


def add_split_slide(prs: Presentation, data: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    dark = data.get("theme") == "dark"
    set_slide_bg(slide, DARK if dark else BG_LIGHT)

    title_color = WHITE if dark else TEXT_DARK
    sub_color = MUTED if dark else MUTED_DARK
    bullet_color = RGBColor(0xE2, 0xE8, 0xF0) if dark else TEXT_DARK

    add_text(slide, Inches(0.75), Inches(0.55), Inches(5.8), Inches(0.7), data["title"], size=34, bold=True, color=title_color)
    add_text(slide, Inches(0.75), Inches(1.25), Inches(5.6), Inches(0.55), data["subtitle"], size=16, color=sub_color)
    add_bullets(slide, Inches(0.75), Inches(2.0), Inches(5.5), Inches(4.5), data["bullets"], size=17, color=bullet_color, spacing=12)

    place_phone(slide, data["image"], Inches(7.0), Inches(0.65), Inches(6.55))


def add_dual_slide(prs: Presentation, data: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK)

    add_text(slide, Inches(0.75), Inches(0.45), Inches(12), Inches(0.65), data["title"], size=34, bold=True, color=WHITE)
    add_text(slide, Inches(0.75), Inches(1.05), Inches(12), Inches(0.5), data["subtitle"], size=17, color=MUTED)

    imgs = data["images"]
    labels = data.get("labels", ["", ""])
    lefts = [Inches(1.6), Inches(7.2)]
    for img, left, label in zip(imgs, lefts, labels):
        place_phone(slide, img, left, Inches(1.55), Inches(5.55))
        if label:
            add_text(slide, left, Inches(6.95), Inches(4.5), Inches(0.4), label, size=15, bold=True, color=RGBColor(0xC7, 0xD2, 0xFE), align=PP_ALIGN.CENTER)


def add_closing_slide(prs: Presentation, data: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, DARK)

    add_text(slide, Inches(0.75), Inches(0.55), Inches(5.8), Inches(0.7), data["title"], size=34, bold=True, color=WHITE)
    add_text(slide, Inches(0.75), Inches(1.2), Inches(5.5), Inches(0.5), data["subtitle"], size=16, color=MUTED)
    add_bullets(slide, Inches(0.75), Inches(1.85), Inches(5.4), Inches(3.2), data["bullets"], size=17, color=RGBColor(0xE2, 0xE8, 0xF0), spacing=11)
    add_text(slide, Inches(0.75), Inches(5.5), Inches(5.5), Inches(0.5), data["contact"], size=20, bold=True, color=RGBColor(0xA5, 0xB4, 0xFC))

    place_phone(slide, data["image"], Inches(7.15), Inches(0.55), Inches(6.6))


def build_pptx() -> Path:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    builders = {
        "cover": add_cover_slide,
        "hero": add_hero_slide,
        "split": add_split_slide,
        "dual": add_dual_slide,
        "closing": add_closing_slide,
    }
    for data in SLIDES:
        builders[data["layout"]](prs, data)

    out = OUT / "koshel-tbank-pitch.pptx"
    prs.save(out)
    return out


class PitchPDF(FPDF):
    def footer(self):
        self.set_y(-10)
        self.set_font("ArialUni", "", 9)
        self.set_text_color(140, 140, 140)
        self.cell(0, 8, str(self.page_no()), align="C")


def _pdf_header(pdf: FPDF, title: str, subtitle: str = "") -> None:
    pdf.set_fill_color(15, 23, 42)
    pdf.rect(0, 0, 297, 22, "F")
    pdf.set_xy(12, 6)
    pdf.set_font("ArialUni", "B", 16)
    pdf.set_text_color(255, 255, 255)
    pdf.cell(0, 8, title, new_x="LMARGIN", new_y="NEXT")
    if subtitle:
        pdf.set_x(12)
        pdf.set_font("ArialUni", "", 10)
        pdf.set_text_color(148, 163, 184)
        pdf.cell(0, 6, subtitle, new_x="LMARGIN", new_y="NEXT")
    pdf.set_y(28)


def build_pdf() -> Path:
    font = "/Library/Fonts/Arial Unicode.ttf"
    pdf = PitchPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(auto=False)
    pdf.add_font("ArialUni", "", font)
    pdf.add_font("ArialUni", "B", font)
    pdf.add_font("ArialUni", "I", font)

    for i, data in enumerate(SLIDES):
        pdf.add_page()
        layout = data["layout"]
        title = data["title"]
        subtitle = data.get("subtitle", data.get("tagline", ""))
        _pdf_header(pdf, title, subtitle)

        if layout == "cover":
            imgs = data["images"]
            xs = [95, 148, 201]
            hs = [115, 105, 95]
            ys = [32, 38, 44]
            for img, x, y, h in zip(imgs, xs, ys, hs):
                path = prepare_phone_image(SCREENSHOTS / img, max_h=900)
                pdf.image(str(path), x=x, y=y, h=h)
            pdf.set_xy(12, 95)
            pdf.set_font("ArialUni", "", 12)
            pdf.set_text_color(165, 180, 252)
            pdf.cell(0, 8, data.get("tagline", ""))

        elif layout == "hero":
            path = prepare_phone_image(SCREENSHOTS / data["image"], max_h=1100)
            pdf.image(str(path), x=78, y=30, h=155)

        elif layout == "dual":
            imgs = data["images"]
            labels = data.get("labels", ["", ""])
            for img, x, label in zip(imgs, [25, 155], labels):
                path = prepare_phone_image(SCREENSHOTS / img, max_h=1000)
                pdf.image(str(path), x=x, y=32, h=140)
                pdf.set_xy(x, 176)
                pdf.set_font("ArialUni", "B", 10)
                pdf.set_text_color(199, 210, 254)
                pdf.cell(80, 6, label, align="C")

        elif layout in ("split", "closing"):
            pdf.set_xy(12, 32)
            pdf.set_font("ArialUni", "", 11)
            pdf.set_text_color(30, 41, 59)
            if data.get("theme") == "dark" or layout == "closing":
                pdf.set_text_color(226, 232, 240)
            for bullet in data.get("bullets", []):
                pdf.multi_cell(88, 6, f"  -  {bullet}")
                pdf.ln(1)
            if data.get("contact"):
                pdf.ln(4)
                pdf.set_font("ArialUni", "B", 13)
                pdf.set_text_color(165, 180, 252)
                pdf.cell(0, 8, data["contact"])
            path = prepare_phone_image(SCREENSHOTS / data["image"], max_h=1050)
            pdf.image(str(path), x=118, y=30, h=150)

    out = OUT / "koshel-tbank-pitch.pdf"
    pdf.output(str(out))
    return out


def main() -> None:
    copy_screenshots()
    pptx = build_pptx()
    pdf = build_pdf()
    print(f"PPTX: {pptx}")
    print(f"PDF:  {pdf}")


if __name__ == "__main__":
    main()
